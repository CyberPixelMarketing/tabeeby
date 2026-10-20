# Builds the starting content.xlsx for the Tabeeby website.
# You do NOT need to run this again: from now on edit content.xlsx directly.
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "content.xlsx"

NAVY, MINT, PALE = "2A428A", "45C79B", "EAF2F1"
HEAD_FONT = Font(name="Cairo", bold=True, color="FFFFFF", size=11)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
KEY_FILL = PatternFill("solid", fgColor=PALE)
THIN = Side(style="thin", color="C9D6DA")
WRAP = Alignment(wrap_text=True, vertical="top")

SETTINGS = [
    # key, English, Arabic, help
    ("NAME", "Tabeeby Medical Center", "مركز طبيبي الطبي", "Center name / اسم المركز"),
    ("PHONE", "050 390 9206", "", "Main phone shown on the site / رقم الاتصال الرئيسي"),
    ("PHONE_2", "06 539 4439", "", "Second phone (landline). Leave empty to hide / رقم ثاني، اتركه فارغاً للإخفاء"),
    ("WHATSAPP", "971503909206", "", "WhatsApp: digits only with country code, no + or 00 / أرقام فقط مع رمز الدولة"),
    ("EMAIL", "tabeebymedsharjah@yahoo.com", "", "Email / البريد الإلكتروني"),
    ("ADDRESS", "Sarh Al Emarat Tower, Al Majaz 3, 2nd floor, offices 201/202/203, Sharjah, United Arab Emirates",
     "برج صرح الإمارات، المجاز 3، الطابق الثاني، مكاتب 201/202/203، الشارقة، الإمارات العربية المتحدة", "Address / العنوان"),
    ("AREA", "Al Majaz, Sharjah", "المجاز، الشارقة", "Short location used in the header strip / الموقع المختصر"),
    ("MAP_QUERY", "Sarh Al Emarat Tower, Al Majaz 3, Sharjah", "", "What Google Maps searches for / ما يبحث عنه في خرائط جوجل"),
    ("HOURS", "", "", "Opening hours. Leave empty to hide / ساعات العمل، اتركها فارغة للإخفاء"),
    ("BOOKING_URL", "", "", "Google Calendar booking page link used when a doctor/department has none. See SETUP guide / رابط صفحة الحجز في تقويم جوجل"),
    ("MESSENGER", "", "", "Messenger link e.g. https://m.me/yourpage (empty = hidden)"),
    ("FACEBOOK", "", "", "Facebook page link (empty = hidden)"),
    ("INSTAGRAM", "", "", "Instagram link (empty = hidden)"),
    ("TIKTOK", "", "", "TikTok link (empty = hidden)"),
    ("HERO_EYEBROW", "Tabeeby Medical Center · Al Majaz, Sharjah", "مركز طبيبي الطبي · المجاز، الشارقة", "Small line above the main title"),
    ("HERO_TITLE", "A smile and skin you feel good about", "ابتسامة وبشرة تشعر بالثقة بهما", "Main title on top of the page"),
    ("HERO_SUB", "Dental care for kids and adults, veneers, crowns, whitening and dentures, plus dermatology and cosmetology. Choose your doctor and book a free time that suits you.",
     "علاج أسنان الأطفال والكبار، الفينير والتيجان والتبييض وأطقم الأسنان، إلى جانب الجلدية والتجميل. اختر طبيبك واحجز الوقت المتاح الذي يناسبك.", "Text under the main title"),
    ("HERO_IMAGE", "", "", "Optional background photo for the top of the page, e.g. img/hero.jpg (empty = brand pattern)"),
    ("ABOUT_TITLE", "Tabeeby", "طبيبي", "About section title"),
    ("ABOUT_TEXT", "Tabeeby Medical Center in Al Majaz, Sharjah, looks after your smile and your skin. Our dentists treat children and adults, from first check-ups and fillings to veneers, crowns, whitening and full dentures, while our dermatology and cosmetology team cares for skin, hair and a fresh, healthy look.\n\nWe keep things simple: you choose the doctor, see their real free times online and book in a few taps. If you need more specialised care, we guide you to the right next step.",
     "يعتني مركز طبيبي الطبي في المجاز بالشارقة بابتسامتك وبشرتك. يعالج أطباء الأسنان لدينا الأطفال والكبار، من الفحص الأول والحشوات إلى الفينير والتيجان والتبييض وأطقم الأسنان الكاملة، بينما يهتم فريق الجلدية والتجميل بالبشرة والشعر وإطلالتك الصحية المشرقة.\n\nنجعل الأمر بسيطاً: تختار الطبيب، وترى أوقاته المتاحة فعلياً، وتحجز بضغطات قليلة. وإذا احتجت رعاية تخصصية أكثر، نرشدك إلى الخطوة الصحيحة.",
     "About text. An empty line starts a new paragraph / نص قسم من نحن"),
    ("ABOUT_IMAGE", "img/post-whitening.png", "", "About section image"),
    ("COLOR_DEEP", "#22397A", "", "Dark navy (dark sections)"),
    ("COLOR_DEEP_SOFT", "#2A428A", "", "Logo navy"),
    ("COLOR_LIGHT", "#EAF2F1", "", "Light background"),
    ("COLOR_LIGHT_SOFT", "#F6FAFA", "", "Lighter background (cards)"),
    ("COLOR_ACCENT", "#45C79B", "", "Logo mint green (accents on dark)"),
    ("COLOR_INK", "#1C2B55", "", "Text colour"),
]

DEPARTMENTS = [
    # id, name_en, name_ar, desc_en, desc_ar, icon
    ("dental", "Dental Services", "خدمات الأسنان",
     "Complete dental care for children and adults: check-ups, cleaning and fillings, veneers, crowns, teeth whitening and full dentures.",
     "رعاية متكاملة لأسنان الأطفال والكبار: الفحص والتنظيف والحشوات، والفينير، والتيجان، وتبييض الأسنان، وأطقم الأسنان الكاملة.", "tooth"),
    ("dermatology", "Dermatology & Cosmetology", "الجلدية والتجميل",
     "Medical skin care and cosmetic treatments for acne, pigmentation, hair loss, eczema and skin rejuvenation.",
     "علاجات جلدية وتجميلية لحب الشباب والتصبغات وتساقط الشعر والإكزيما ونضارة البشرة.", "skin"),
]

DOCTORS = [
    # id, name_en, name_ar, title_en, title_ar, department
    ("taima-tanjour", "Dr. Taima Tanjour", "د. تيماء طنجور", "General Dentist", "طبيبة أسنان عامة", "dental"),
    ("lora-el-hourani", "Dr. Lora El Hourani", "د. لورا الحوراني", "General Dentist", "طبيبة أسنان عامة", "dental"),
    ("amany-zoarob", "Dr. Amany Zoarob", "د. أماني زعرب", "General Dentist", "طبيبة أسنان عامة", "dental"),
    ("nadia-mohamed", "Dr. Nadia Mohamed", "د. نادية محمد", "Specialist Dermatology", "أخصائية الأمراض الجلدية", "dermatology"),
    ("samar-matter", "Dr. Samar Matter", "د. سمر مطر", "Specialist Dermatology & Cosmetology", "أخصائية الجلدية والتجميل", "dermatology"),
]

FEATURES = [
    ("Dental care for every age", "أسنان لكل الأعمار", "From a child's first check-up to veneers, crowns and full dentures for adults.", "من أول فحص للطفل إلى الفينير والتيجان وأطقم الأسنان الكاملة للكبار."),
    ("Real-time booking", "حجز فوري", "See your doctor's actual free times and confirm your visit in a few taps.", "شاهد الأوقات المتاحة فعلياً لطبيبك وأكّد زيارتك بضغطات قليلة."),
    ("Skin and beauty specialists", "أخصائيات الجلدية والتجميل", "Medical dermatology and cosmetic treatments planned around your skin.", "علاجات جلدية طبية وتجميلية مصممة حسب طبيعة بشرتك."),
    ("In the heart of Al Majaz", "في قلب المجاز", "Easy to reach in Sarh Al Emarat Tower, and a WhatsApp message away.", "سهل الوصول في برج صرح الإمارات، وعلى بعد رسالة واتساب."),
]

SERVICES = [
    # department, name_en, name_ar, text_en, text_ar, icon
    ("dental", "Kids' dentistry", "طب أسنان الأطفال",
     "Gentle check-ups, fluoride, fissure sealants and fillings that help children feel at ease in the chair.",
     "فحوصات لطيفة وفلورايد وسدّ الشقوق وحشوات تجعل الطفل مرتاحاً على كرسي الأسنان.", "child"),
    ("dental", "Adult dentistry", "طب أسنان الكبار",
     "Check-ups, scaling and polishing, tooth-coloured fillings and treatment of tooth pain and gum problems.",
     "الفحص الدوري وتنظيف الجير والتلميع والحشوات التجميلية وعلاج آلام الأسنان ومشاكل اللثة.", "tooth"),
    ("dental", "Veneers", "الفينير",
     "Thin custom shells that fix the shape, size and colour of front teeth for a natural, even smile.",
     "قشور رقيقة مصممة لك تصحح شكل الأسنان الأمامية وحجمها ولونها لابتسامة طبيعية ومتناسقة.", "sparkle"),
    ("dental", "Crowns", "التيجان",
     "Strong, natural-looking crowns that protect and restore broken, weak or root-treated teeth.",
     "تيجان قوية بمظهر طبيعي تحمي الأسنان المكسورة أو الضعيفة أو المعالجة عصبياً وتعيد بناءها.", "crown"),
    ("dental", "Teeth whitening", "تبييض الأسنان",
     "Safe in-clinic whitening that lifts stains from coffee, tea and smoking, with aftercare advice so it lasts.",
     "تبييض آمن في العيادة يزيل تصبغات القهوة والشاي والتدخين، مع إرشادات بعد التبييض لتدوم النتيجة.", "sparkle"),
    ("dental", "Full dentures", "أطقم الأسنان الكاملة",
     "Comfortable, well-fitted complete dentures that bring back chewing, speech and a confident smile.",
     "أطقم كاملة مريحة ومضبوطة المقاس تعيد لك المضغ والنطق والابتسامة الواثقة.", "smile"),
    ("dermatology", "Acne & acne scars", "حب الشباب وآثاره",
     "Treatment plans that calm active acne and reduce marks and scars.",
     "خطط علاج تهدّئ حب الشباب النشط وتخفف البقع والندبات.", "skin"),
    ("dermatology", "Pigmentation & melasma", "التصبغات والكلف",
     "Even out dark spots, melasma and uneven skin tone with medical-grade treatments.",
     "توحيد لون البشرة وعلاج البقع الداكنة والكلف بعلاجات طبية.", "drop"),
    ("dermatology", "Eczema, psoriasis & allergies", "الإكزيما والصدفية والحساسية",
     "Diagnosis and long-term care for itchy, inflamed and sensitive skin.",
     "تشخيص ومتابعة طويلة الأمد للبشرة الحساسة والملتهبة والمصابة بالحكة.", "shield"),
    ("dermatology", "Hair loss & scalp care", "تساقط الشعر وفروة الرأس",
     "Find the cause of hair thinning and start a treatment plan that fits you.",
     "معرفة سبب تساقط الشعر وبدء خطة علاج تناسبك.", "hair"),
    ("dermatology", "Skin rejuvenation & peels", "نضارة البشرة والتقشير",
     "Chemical peels and skin-boosting treatments for a fresher, brighter complexion.",
     "التقشير الكيميائي وعلاجات تحفيز البشرة لإطلالة أكثر نضارة وإشراقاً.", "sparkle"),
    ("dermatology", "Anti-ageing treatments", "علاجات مكافحة الشيخوخة",
     "Softening fine lines and restoring volume with cosmetic injectables, done by a specialist.",
     "تخفيف الخطوط الدقيقة واستعادة الامتلاء بالحقن التجميلية على يد أخصائية.", "face"),
]

FAQ = [
    ("How do I book an appointment?", "كيف أحجز موعداً؟",
     "Choose a department and a doctor, then pick any free time shown in the calendar and enter your name and phone. You receive a confirmation by email straight away.",
     "اختر القسم والطبيب، ثم اختر أي وقت متاح في التقويم وأدخل اسمك ورقم هاتفك. تصلك رسالة تأكيد على بريدك الإلكتروني مباشرة."),
    ("Can I change or cancel my appointment?", "هل يمكنني تغيير الموعد أو إلغاؤه؟",
     "Yes. Use the link in your confirmation email, or message us on WhatsApp and we will move it for you.",
     "نعم. استخدم الرابط الموجود في رسالة التأكيد، أو راسلنا على واتساب وسنغيّر الموعد لك."),
    ("Do you treat children's teeth?", "هل تعالجون أسنان الأطفال؟",
     "Yes. Our dentists see children and adults. Book as usual and write the child's name and age in the notes.",
     "نعم. أطباء الأسنان لدينا يعالجون الأطفال والكبار. احجز كالمعتاد واكتب اسم الطفل وعمره في الملاحظات."),
    ("How long does teeth whitening last?", "كم تدوم نتيجة تبييض الأسنان؟",
     "With good care it lasts a long time. For the first days after whitening avoid coffee, tea, coloured food and drinks, and smoking.",
     "مع العناية الجيدة تدوم طويلاً. في الأيام الأولى بعد التبييض تجنّب القهوة والشاي والأطعمة والمشروبات الملونة والتدخين."),
    ("Do you accept my insurance?", "هل تقبلون التأمين الخاص بي؟",
     "Send us your insurance card details on WhatsApp before your visit and we will confirm your coverage.",
     "أرسل لنا بيانات بطاقة التأمين على واتساب قبل زيارتك وسنؤكد لك التغطية."),
    ("Where exactly are you?", "أين يقع المركز بالتحديد؟",
     "Sarh Al Emarat Tower, Al Majaz 3, Sharjah, on the 2nd floor in offices 201, 202 and 203.",
     "برج صرح الإمارات، المجاز 3، الشارقة، في الطابق الثاني مكاتب 201 و202 و203."),
]

HOW_TO = [
    ("HOW TO MANAGE THE WEBSITE WITH THIS FILE", "طريقة إدارة الموقع من هذا الملف"),
    ("", ""),
    ("1. Each tab controls one part of the website. Edit the cells, save the file, and upload it to the website folder with the same name: content.xlsx", "1. كل ورقة تتحكم بجزء من الموقع. عدّل الخلايا واحفظ الملف ثم ارفعه إلى مجلد الموقع بنفس الاسم: content.xlsx"),
    ("2. Settings: phone, WhatsApp, email, address, texts, colours and the Google Calendar booking link.", "2. Settings: الهاتف والواتساب والبريد والعنوان والنصوص والألوان ورابط الحجز في تقويم جوجل."),
    ("3. Departments: add a row to add a department. The id must be short, English, no spaces (e.g. ent). Set show to no to hide it.", "3. Departments: أضف صفاً لإضافة قسم. الـ id قصير بالإنجليزية بدون مسافات. اكتب no في عمود show للإخفاء."),
    ("4. Doctors: add a row per doctor. department must match an id from the Departments tab. photo = image path (e.g. img/doctors/reham.jpg) or a full https link.", "4. Doctors: صف لكل طبيب. عمود department يجب أن يطابق id من ورقة الأقسام. photo = مسار الصورة أو رابط https كامل."),
    ("5. booking_url (Doctors / Departments): paste that doctor's Google Calendar booking page link. Empty = uses the department link, then BOOKING_URL in Settings.", "5. booking_url: الصق رابط صفحة الحجز الخاصة بالطبيب من تقويم جوجل. إن تُرك فارغاً يُستخدم رابط القسم ثم الرابط العام في Settings."),
    ("6. order: lower numbers appear first. Do not rename the tabs or the column titles in row 1.", "6. order: الأرقام الأصغر تظهر أولاً. لا تغيّر أسماء الأوراق ولا عناوين الأعمدة في الصف الأول."),
    ("7. Services tab = the treatments shown inside each department. Icons available: tooth, child, smile, crown, sparkle, skin, drop, shield, hair, face, stethoscope, heart", "7. ورقة Services = العلاجات المعروضة داخل كل قسم. الأيقونات المتاحة: tooth, child, smile, crown, sparkle, skin, drop, shield, hair, face, stethoscope, heart"),
]


def style_header(ws, widths):
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "B2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            c.border = Border(bottom=THIN)
            c.font = Font(name="Cairo", size=10)


def yes_no(ws, col, n):
    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{n + 200}")


wb = Workbook()
ws = wb.active
ws.title = "How to use"
ws.sheet_view.showGridLines = False
for r, (en, ar) in enumerate(HOW_TO, 1):
    ws.cell(row=r, column=1, value=en).alignment = WRAP
    c = ws.cell(row=r, column=2, value=ar)
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
    for col in (1, 2):
        ws.cell(row=r, column=col).font = Font(name="Cairo", size=11, bold=(r == 1), color=(NAVY if r == 1 else "1C2B55"))
ws.column_dimensions["A"].width = 80
ws.column_dimensions["B"].width = 80

ws = wb.create_sheet("Settings")
ws.append(["key", "english", "arabic", "help"])
for row in SETTINGS:
    ws.append(list(row))
style_header(ws, [18, 60, 60, 60])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=1).fill = KEY_FILL
    ws.cell(row=r, column=1).font = Font(name="Cairo", size=10, bold=True)
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)

ws = wb.create_sheet("Departments")
ws.append(["id", "name_en", "name_ar", "description_en", "description_ar", "icon", "image", "booking_url", "show", "order"])
for i, d in enumerate(DEPARTMENTS, 1):
    ws.append(list(d) + ["", "", "yes", i])
style_header(ws, [18, 26, 22, 50, 50, 12, 24, 40, 8, 8])
yes_no(ws, "I", len(DEPARTMENTS))

ws = wb.create_sheet("Doctors")
ws.append(["id", "name_en", "name_ar", "title_en", "title_ar", "department", "photo", "languages_en", "languages_ar", "bio_en", "bio_ar", "booking_url", "show", "order"])
for i, d in enumerate(DOCTORS, 1):
    ws.append(list(d) + ["", "", "", "", "", "", "yes", i])
style_header(ws, [18, 30, 26, 34, 32, 18, 30, 20, 20, 40, 40, 40, 8, 8])
yes_no(ws, "M", len(DOCTORS))
dv = DataValidation(type="list", formula1="=Departments!$A$2:$A$100", allow_blank=False)
ws.add_data_validation(dv)
dv.add("F2:F300")

ws = wb.create_sheet("Features")
ws.append(["title_en", "title_ar", "text_en", "text_ar", "show"])
for f in FEATURES:
    ws.append(list(f) + ["yes"])
style_header(ws, [28, 26, 60, 60, 8])
yes_no(ws, "E", len(FEATURES))

ws = wb.create_sheet("Services")
ws.append(["department", "name_en", "name_ar", "text_en", "text_ar", "icon", "image", "show", "order"])
for i, sv in enumerate(SERVICES, 1):
    ws.append(list(sv) + ["", "yes", i])
style_header(ws, [16, 30, 28, 55, 55, 12, 24, 8, 8])
yes_no(ws, "H", len(SERVICES))

ws = wb.create_sheet("FAQ")
ws.append(["question_en", "question_ar", "answer_en", "answer_ar", "show"])
for f in FAQ:
    ws.append(list(f) + ["yes"])
style_header(ws, [40, 40, 60, 60, 8])
yes_no(ws, "E", len(FAQ))

wb.save(OUT)
print("wrote", OUT)
