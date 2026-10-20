# Optional: copies content.xlsx into content-fallback.js, which the site only uses
# when content.xlsx cannot be loaded (for example when index.html is opened
# straight from the computer instead of a web server).
#   python3 tools/xlsx_to_fallback.py
import json, os, sys
from openpyxl import load_workbook

root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(root, "content.xlsx")
wb = load_workbook(src, data_only=True)
out = {}
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    if not rows or ws.title == "How to use":
        continue
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    items = []
    for r in rows[1:]:
        if not any(v not in (None, "") for v in r):
            continue
        items.append({h: ("" if v is None else v) for h, v in zip(head, r) if h})
    out[ws.title] = items
with open(os.path.join(root, "content-fallback.js"), "w", encoding="utf-8") as f:
    f.write("/* Generated from content.xlsx — edit the Excel file, not this one. */\n")
    f.write("window.CONTENT_FALLBACK = " + json.dumps(out, ensure_ascii=False, indent=1) + ";\n")
print("wrote content-fallback.js:", {k: len(v) for k, v in out.items()})
